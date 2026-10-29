#!/usr/bin/env python3
"""
Builds the Word version of the roles annexure.

Takes the content straight from build_roles.py, so the Word file and the
PDF never drift apart: everything above the `def rows(` line in that file
is data, so that part is executed and the rendering code is not.

Run with python3.13, which is where python-docx is installed.
"""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "/home/user/personal_v1/mahapar_review_assets/build_roles.py"
OUT = "/home/user/personal_v1/MahaPAR_APAR_roles_and_responsibilities.docx"

NAVY = RGBColor(0x1B, 0x3A, 0x6B)
FILL = "CBD9EC"

src = open(SRC).read()
data = {}
exec(src[: src.index("def rows(")], data)          # data only; no rendering, no chromium
CHAIN, ADMIN = data["CHAIN"], data["ADMIN"]
ESCALATION, TIMELINE = data["ESCALATION"], data["TIMELINE"]

ENT = {"&ndash;": "–", "&minus;": "−", "&rsquo;": "’",
       "&lsquo;": "‘", "&amp;": "&", "&#8505;&#65039;": ""}


def clean(t):
    for k, v in ENT.items():
        t = t.replace(k, v)
    return t


def runs(para, html, base_italic=False):
    """Writes html with <b> and <i> as real Word runs; <ul> becomes dashes."""
    html = clean(html)
    for chunk in re.split(r"(<b>.*?</b>|<i>.*?</i>)", html, flags=re.S):
        if not chunk:
            continue
        bold = chunk.startswith("<b>")
        ital = chunk.startswith("<i>") or base_italic
        text = re.sub(r"</?[bi]>", "", chunk)
        text = re.sub(r"[ \t]*<ul>\s*", "\n", text)
        text = re.sub(r"\s*</ul>[ \t]*", "", text)
        text = re.sub(r"[ \t]*<li>[ \t]*", "– ", text)
        text = re.sub(r"[ \t]*</li>[ \t]*\n?", "\n", text)
        text = re.sub(r"<[^>]+>", "", text)
        text = re.sub(r"[ \t]+", " ", text).strip("\n")
        for i, line in enumerate(text.split("\n")):
            if i:
                para.add_run().add_break()
            r = para.add_run(line)
            r.bold, r.italic = bold, ital


def shade(cell, hex_fill):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(el)


def fixed_layout(table, widths):
    """Word only honours cell widths when the table layout is fixed."""
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    grid = OxmlElement("w:tblGrid")
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(w.twips)))
        grid.append(gc)
    table._tbl.insert(1, grid)
    old = table._tbl.find(qn("w:tblGrid"))
    for g in table._tbl.findall(qn("w:tblGrid"))[1:]:
        table._tbl.remove(g)


def borders(table):
    tbl = table._tbl.tblPr
    el = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), "000000")
        el.append(e)
    tbl.append(el)


doc = Document()
st = doc.styles["Normal"]
st.font.name, st.font.size = "Calibri", Pt(10)
st.paragraph_format.space_after = Pt(4)

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.5)
sec.top_margin = sec.bottom_margin = Cm(1.8)
USABLE = Cm(18)


def heading(text, size, space_before=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(clean(text))
    r.bold, r.font.size, r.font.color.rgb = True, Pt(size), NAVY
    return p


def para(text, italic=False, center=False, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    runs(p, text, base_italic=italic)
    return p


def table(headers, widths, rows_data, header_fill=True):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    borders(t)
    fixed_layout(t, widths)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.width = widths[i]
        if header_fill:
            shade(c, FILL)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.bold, r.font.color.rgb = True, NAVY
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row in rows_data:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].width = widths[i]
            runs(cells[i].paragraphs[0], val)
    return t


# ---- title and opening box -------------------------------------------------
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("APAR under MahaPAR: who does what, by when, and how")
r.bold, r.font.size, r.font.color.rgb = True, Pt(16), NAVY
para("Draft for discussion with GAD | IDinsight | Based on Form A, for Group A and "
     "Group B officials, current draft", center=True)

lead = doc.add_table(rows=1, cols=1)
lead.autofit = False
borders(lead)
lead.rows[0].cells[0].width = USABLE
box = lead.rows[0].cells[0]
LEAD = [
 "The APAR covers the financial year just ended, 1 April to 31 March, and the cycle opens "
 "on the following 1 April. The APAR for 2025-26 opens on 1 April 2026.",
 "Every part is filed in MahaPAR and signed with a digital signature or eSign. Paper forms "
 "are not accepted.",
 "If a stage is not completed by its date, MahaPAR forwards the APAR to the next stage, "
 "records the delay, and deducts timeliness marks from the APAR of the official who held it. "
 "These marks are 5 per cent of the final score, that is &minus;0.5 to +0.5 on the 10-point "
 "scale, and they sit in Part 6, which MahaPAR generates on its own. Timeliness is counted "
 "for every role an official holds in the cycle: their own self-appraisal, their assessments "
 "as a reporting officer, and their reviews as a reviewing officer.",
]
for i, t in enumerate(LEAD):
    p = box.paragraphs[0] if i == 0 else box.add_paragraph()
    runs(p, t)

# ---- A. the cycle ----------------------------------------------------------
heading("A. The cycle at a glance", 13)
t = table(["Date", "Stage"], [Cm(4), Cm(14)],
          [(f"<b>{d}</b>", e) for d, e in TIMELINE])
para("<b>Note on the dates.</b> The dates shown are the ones proposed in the stakeholder "
     "draft, modelled on Gujarat&rsquo;s cycle, which gives each stage about 30 days. They "
     "will be finalised in consultation with GAD.", space_before=8)
para("After the reviewing stage the completed APAR is disclosed to the official by the "
     "custodian. The official may submit one representation within 15 days of disclosure. "
     "No entries are accepted after the cycle closes.")

# ---- B and C. the role tables ---------------------------------------------
W = [Cm(5.8), Cm(1.8), Cm(2.4), Cm(8.0)]
HEAD = ["Responsibility", "Form", "By when", "How it is done in MahaPAR"]


def role_block(title, scope, items):
    heading(title, 11.5)
    p = para(scope, italic=True)
    p.paragraph_format.keep_with_next = True
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{len(items)} action{'s' if len(items) != 1 else ''}")
    r.bold = True
    table(HEAD, W, [(a, b or "–", c or "–", d or "–") for a, b, c, d in items])


heading("B. Officers in the APAR chain", 13)
for block in CHAIN:
    role_block(*block)

doc.add_page_break()
heading("C. Administrative and monitoring roles", 13, space_before=0)
for block in ADMIN:
    role_block(*block)

# ---- D. what happens if a stage is missed ---------------------------------
heading("D. What happens if a stage is not completed", 13)
table(["Situation", "What happens next"], [Cm(7), Cm(11)], ESCALATION)
para("Follow-up runs upward: official, custodian, Cadre Monitoring Authority, Sub Department "
     "Monitoring Authority, GAD. Reminders to an official are sent by the custodian. Each "
     "level above follows up with the level below it, not with the official.", space_before=8)

doc.save(OUT)
print("wrote", OUT)
